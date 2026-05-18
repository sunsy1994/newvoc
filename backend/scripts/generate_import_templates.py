from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT_DIR = Path(__file__).resolve().parents[2]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from backend.app.template_registry import TEMPLATES


SAMPLE_ROWS: dict[str, list[object]] = {
    "tpl-event": [
        "EVT-2026-001",
        "某品牌SUV高速事故舆情",
        "产品质量",
        "某品牌",
        "X7",
        "2026-01-06 00:00:00",
        "进行中",
        "车辆事故引发多平台讨论，需持续跟踪传播走势。",
        "2026-03-20 00:00:00",
        "事故|高速|安全|售后",
    ],
    "tpl-content": [
        "B站",
        "https://example.com/posts/cnt-1001",
        "某测评号",
        "X7高速事故复盘：问题出在机械还是处置",
        "2026-03-07 10:12:00",
        "CNT-1001",
        "EVT-2026-001",
        "AUTH-301",
        "KOL",
        "是",
        "汽车测评",
        1260000,
        "结合公开视频和用户反馈，对事件经过进行拆解。",
        "视频",
        "横版视频",
        6120,
        1864,
        930,
        1210,
        35600,
    ],
    "tpl-comment": [
        "CNT-1001",
        "普通车主A",
        "这次事故让人担心X7的高速稳定性，建议品牌尽快给解释。",
        "2026-03-07 12:22:00",
        "担忧质量",
        "求证原因",
        "负向",
        "家庭用户",
        "CMT-9001",
        "B站",
        "USR-001",
        "",
        1,
        624,
        115,
        "https://example.com/comments/cmt-9001",
        "观望期",
        "安全优先",
    ],
    "tpl-account": [
        "B站",
        "某测评号",
        "AUTH-301",
        "https://example.com/avatar/auth-301.png",
        1260000,
        126,
        5320000,
        "长期发布汽车测评与事故解析内容。",
        "汽车领域认证创作者",
        "KOL",
        "汽车测评",
        "理性分析型",
        "成熟期",
        '["事故复盘","汽车测评"]',
    ],
}


OUTPUT_DIRS = [
    Path("app/public/data-import-templates"),
    Path("app/dist/data-import-templates"),
]


def build_template(template: dict, output_dir: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = template["source_name"]

    headers = [*template["required_headers"], *template["optional_headers"]]
    sample_row = SAMPLE_ROWS.get(template["id"], [])

    worksheet.append(headers)
    if sample_row:
        worksheet.append(sample_row)

    header_fill = PatternFill(fill_type="solid", fgColor="DCEBFF")
    header_font = Font(bold=True, color="1F2937")

    for index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.column_dimensions[cell.column_letter].width = max(len(str(header)) * 2, 14)

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(output_dir / template["file_name"])


def main() -> None:
    for output_dir in OUTPUT_DIRS:
        for template in TEMPLATES:
            build_template(template, output_dir)
        print(f"Generated templates in {output_dir}")


if __name__ == "__main__":
    main()
