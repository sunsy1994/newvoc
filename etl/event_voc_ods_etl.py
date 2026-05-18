from __future__ import annotations

import argparse
import hashlib
import json
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE_DIR = ROOT_DIR / "templates" / "event-voc"
DEFAULT_SAMPLE_DIR = ROOT_DIR / "samples" / "event_voc_etl_sample"


@dataclass(frozen=True)
class TemplateSpec:
    key: str
    title: str
    file_stem: str
    columns: list[tuple[str, str, bool, str]]
    sample_row: dict[str, Any]


EVENT_SPEC = TemplateSpec(
    key="event",
    title="VOC看事件-事件上传模板",
    file_stem="event_upload_template",
    columns=[
        ("原始事件ID", "raw_event_id", True, "你自己维护的事件编号，建议稳定不重复"),
        ("事件名称", "event_name", True, "事件标题"),
        ("事件类型", "event_type", True, "新品上市/品牌传播/价格权益/产品质量/服务体验/事故舆情/竞品对比/用户口碑/其他"),
        ("品牌名称", "brand_name", False, "品牌"),
        ("车型名称", "model_name", False, "车型"),
        ("事件开始时间", "start_time", False, "YYYY-MM-DD HH:MM:SS"),
        ("事件结束时间", "end_time", False, "YYYY-MM-DD HH:MM:SS，可为空"),
        ("事件状态", "event_status", True, "进行中/已结束/归档"),
        ("事件关键词", "keyword_list", False, "多个关键词用逗号、顿号或竖线分隔"),
        ("事件描述", "event_desc", False, "事件背景说明"),
    ],
    sample_row={
        "原始事件ID": "EVT-2026-001",
        "事件名称": "A车型上市传播",
        "事件类型": "新品上市",
        "品牌名称": "某品牌",
        "车型名称": "A车型",
        "事件开始时间": "2026-05-01 00:00:00",
        "事件结束时间": "",
        "事件状态": "进行中",
        "事件关键词": "A车型,上市,权益",
        "事件描述": "上市传播阶段重点观察用户对价格和配置的反馈",
    },
)


CONTENT_SPEC = TemplateSpec(
    key="content",
    title="VOC看事件-内容上传模板",
    file_stem="content_upload_template",
    columns=[
        ("原始内容ID", "raw_content_id", False, "有平台内容ID就填，没有可为空"),
        ("原始事件ID", "raw_event_id", True, "填写事件上传模板里的原始事件ID"),
        ("标准事件ID", "event_id", False, "高级字段，普通上传可不填"),
        ("平台", "platform", True, "抖音/快手/小红书/微博/B站/懂车帝/汽车之家等"),
        ("原始链接", "source_url", True, "内容原始链接，用于去重"),
        ("内容标题", "title", True, "标题；无标题时可用正文前若干字"),
        ("内容正文", "content_text", False, "正文"),
        ("内容类型", "content_type", False, "官方内容/媒体内容/达人内容/用户内容/经销商内容"),
        ("媒介形态", "media_form", False, "视频/图文/直播/其他"),
        ("发布时间", "published_at", True, "YYYY-MM-DD HH:MM:SS"),
        ("点赞数", "like_cnt", False, "空值按0处理"),
        ("评论数", "comment_cnt", False, "平台显示评论数，空值按0处理"),
        ("分享数", "share_cnt", False, "空值按0处理"),
        ("收藏数", "favorite_cnt", False, "空值按0处理"),
        ("播放/阅读数", "view_cnt", False, "不计入互动量"),
        ("原始作者ID", "raw_author_id", False, "有平台作者ID就填，没有可为空"),
        ("作者名称", "author_name", True, "作者昵称/账号名"),
        ("作者主页链接", "author_home_url", False, "作者主页，用于作者去重"),
        ("作者类型", "author_type", False, "官方号/媒体号/达人/经销商/普通用户/其他"),
        ("是否KOL", "is_kol", False, "是/否，或 true/false"),
        ("作者粉丝数", "fans_cnt", False, "空值按0处理"),
        ("作者简介", "author_desc", False, "作者简介"),
    ],
    sample_row={
        "原始内容ID": "",
        "原始事件ID": "EVT-2026-001",
        "标准事件ID": "",
        "平台": "抖音",
        "原始链接": "https://example.com/post/a-001",
        "内容标题": "A车型上市，价格到底香不香？",
        "内容正文": "这次上市主要看价格和配置，评论区讨论很热。",
        "内容类型": "达人内容",
        "媒介形态": "视频",
        "发布时间": "2026-05-01 20:30:00",
        "点赞数": 1200,
        "评论数": 318,
        "分享数": 88,
        "收藏数": 56,
        "播放/阅读数": 50213,
        "原始作者ID": "",
        "作者名称": "车圈老张",
        "作者主页链接": "https://example.com/user/laozhang",
        "作者类型": "达人",
        "是否KOL": "是",
        "作者粉丝数": 120000,
        "作者简介": "汽车测评博主",
    },
)


COMMENT_SPEC = TemplateSpec(
    key="comment",
    title="VOC看事件-评论上传模板",
    file_stem="comment_upload_template",
    columns=[
        ("原始评论ID", "raw_comment_id", False, "有平台评论ID就填，没有可为空"),
        ("所属内容ID", "content_id", False, "可填系统content_id或内容模板里的原始内容ID"),
        ("所属内容原始链接", "content_source_url", True, "推荐填写内容原始链接，用于匹配帖子"),
        ("平台", "platform", False, "为空时可从内容继承"),
        ("评论作者ID", "comment_author_id", False, "第一版不做跨平台用户识别"),
        ("评论作者昵称", "comment_author_name", True, "评论者昵称"),
        ("父评论ID", "parent_comment_id", False, "回复关系，可为空"),
        ("回复层级", "reply_level", False, "空值按1处理"),
        ("评论正文", "comment_text", True, "VOC原文"),
        ("评论发布时间", "published_at", True, "YYYY-MM-DD HH:MM:SS"),
        ("评论点赞数", "like_cnt", False, "空值按0处理"),
        ("评论回复数", "reply_cnt", False, "空值按0处理"),
        ("评论原始链接", "source_url", False, "评论链接，可为空"),
    ],
    sample_row={
        "原始评论ID": "",
        "所属内容ID": "",
        "所属内容原始链接": "https://example.com/post/a-001",
        "平台": "抖音",
        "评论作者ID": "",
        "评论作者昵称": "喜欢旅行的小王",
        "父评论ID": "",
        "回复层级": 1,
        "评论正文": "这个价格如果有置换补贴就很香，想看看配置差异。",
        "评论发布时间": "2026-05-01 21:10:00",
        "评论点赞数": 23,
        "评论回复数": 4,
        "评论原始链接": "https://example.com/comment/c-001",
    },
)


SPECS = [EVENT_SPEC, CONTENT_SPEC, COMMENT_SPEC]


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def stable_hash(*parts: Any) -> str:
    text = "||".join(normalize_text(part).lower() for part in parts)
    return hashlib.sha1(text.encode("utf-8")).hexdigest()[:16]


def to_bool(value: Any) -> bool:
    text = normalize_text(value).lower()
    return text in {"1", "true", "yes", "y", "是", "有", "kol"}


def to_int(value: Any, default: int = 0) -> int:
    text = normalize_text(value)
    if not text:
        return default
    try:
        return int(float(text))
    except ValueError:
        return default


def clean_optional(value: Any) -> Any:
    text = normalize_text(value)
    return text if text else None


def parse_time(value: Any) -> pd.Timestamp | pd.NaT:
    if value is None or normalize_text(value) == "":
        return pd.NaT
    return pd.to_datetime(value, errors="coerce")


def spec_column_map(spec: TemplateSpec) -> dict[str, str]:
    return {label: field for label, field, _required, _desc in spec.columns}


def required_fields(spec: TemplateSpec) -> list[str]:
    return [field for _label, field, required, _desc in spec.columns if required]


def write_template(spec: TemplateSpec, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    headers = [label for label, _field, _required, _desc in spec.columns]
    desc_row = {label: desc for label, _field, _required, desc in spec.columns}
    required_row = {label: "必填" if required else "选填" for label, _field, required, _desc in spec.columns}
    dataframe = pd.DataFrame([required_row, desc_row, spec.sample_row], columns=headers)
    dataframe.to_csv(output_dir / f"{spec.file_stem}.csv", index=False, encoding="utf-8-sig")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = spec.title[:31]
    worksheet.append(headers)
    worksheet.append([required_row.get(header, "") for header in headers])
    worksheet.append([desc_row.get(header, "") for header in headers])
    worksheet.append([spec.sample_row.get(header, "") for header in headers])

    header_fill = PatternFill(fill_type="solid", fgColor="DCEBFF")
    meta_fill = PatternFill(fill_type="solid", fgColor="F5F7FA")
    header_font = Font(bold=True, color="1F2937")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2, max_row=3):
        for cell in row:
            cell.fill = meta_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in worksheet.iter_rows(min_row=4, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, header in enumerate(headers, start=1):
        column_letter = worksheet.cell(row=1, column=index).column_letter
        worksheet.column_dimensions[column_letter].width = min(max(len(header) * 2, 16), 28)
    worksheet.freeze_panes = "A4"
    workbook.save(output_dir / f"{spec.file_stem}.xlsx")


def generate_templates(output_dir: Path) -> None:
    for spec in SPECS:
        write_template(spec, output_dir)


def write_sample_inputs(input_dir: Path) -> None:
    input_dir.mkdir(parents=True, exist_ok=True)
    for spec in SPECS:
        headers = [label for label, _field, _required, _desc in spec.columns]
        pd.DataFrame([spec.sample_row], columns=headers).to_excel(
            input_dir / f"{spec.key}_upload.xlsx",
            index=False,
        )


def read_upload(input_dir: Path, spec: TemplateSpec) -> pd.DataFrame:
    candidates = [
        input_dir / f"{spec.key}_upload.xlsx",
        input_dir / f"{spec.key}_upload.csv",
        input_dir / f"{spec.file_stem}.xlsx",
        input_dir / f"{spec.file_stem}.csv",
    ]
    path = next((candidate for candidate in candidates if candidate.exists()), None)
    if path is None:
        return pd.DataFrame(columns=[field for _label, field, _required, _desc in spec.columns])
    if path.suffix.lower() == ".csv":
        dataframe = pd.read_csv(path)
    else:
        dataframe = pd.read_excel(path)
    dataframe = dataframe.dropna(how="all")
    dataframe = dataframe.rename(columns=spec_column_map(spec))
    field_names = [field for _label, field, _required, _desc in spec.columns]
    for field in field_names:
        if field not in dataframe.columns:
            dataframe[field] = None
    dataframe = dataframe[field_names]
    return dataframe


def validate_required(dataframe: pd.DataFrame, spec: TemplateSpec) -> list[str]:
    errors: list[str] = []
    for field in required_fields(spec):
        missing_rows = dataframe.index[
            dataframe[field].isna() | (dataframe[field].astype(str).str.strip() == "")
        ].tolist()
        if missing_rows:
            row_nums = ", ".join(str(row + 2) for row in missing_rows[:10])
            errors.append(f"{spec.title} 字段 {field} 缺失，行号：{row_nums}")
    return errors


def standardize_events(ods_event: pd.DataFrame, ingest_batch_id: str) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for row in ods_event.to_dict(orient="records"):
        raw_event_id = clean_optional(row.get("raw_event_id"))
        event_key = raw_event_id or f"event_{stable_hash(row.get('event_name'), row.get('brand_name'), row.get('model_name'), row.get('start_time'))}"
        event_id = raw_event_id or event_key
        rows.append(
            {
                "event_id": event_id,
                "event_key": event_key,
                "event_name": clean_optional(row.get("event_name")),
                "event_type": clean_optional(row.get("event_type")) or "其他",
                "brand_name": clean_optional(row.get("brand_name")),
                "model_name": clean_optional(row.get("model_name")),
                "start_time": parse_time(row.get("start_time")),
                "end_time": parse_time(row.get("end_time")),
                "event_status": clean_optional(row.get("event_status")) or "进行中",
                "keyword_list": clean_optional(row.get("keyword_list")),
                "event_desc": clean_optional(row.get("event_desc")),
                "ingest_batch_id": ingest_batch_id,
                "raw_source_key": "event_upload",
            }
        )
    return pd.DataFrame(rows).drop_duplicates("event_key", keep="last")


def standardize_content(
    ods_content: pd.DataFrame,
    dwd_event: pd.DataFrame,
    ingest_batch_id: str,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    event_by_id = {row["event_id"]: row["event_id"] for row in dwd_event.to_dict(orient="records")}
    event_by_key = {row["event_key"]: row["event_id"] for row in dwd_event.to_dict(orient="records")}
    authors: list[dict[str, Any]] = []
    contents: list[dict[str, Any]] = []
    rel_event_content: list[dict[str, Any]] = []
    rel_author_content: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []

    for row in ods_content.to_dict(orient="records"):
        raw_event_id = clean_optional(row.get("raw_event_id"))
        input_event_id = clean_optional(row.get("event_id"))
        event_id = None
        if input_event_id:
            event_id = event_by_id.get(input_event_id)
        if event_id is None and raw_event_id:
            event_id = event_by_id.get(raw_event_id) or event_by_key.get(raw_event_id)
        if event_id is None:
            rejected.append({"reason": "无法匹配事件", **row})
            continue

        platform = clean_optional(row.get("platform")) or ""
        source_url = clean_optional(row.get("source_url")) or ""
        raw_content_id = clean_optional(row.get("raw_content_id"))
        content_key = f"content_{stable_hash(platform, source_url)}"
        content_id = raw_content_id or content_key

        raw_author_id = clean_optional(row.get("raw_author_id"))
        author_home_url = clean_optional(row.get("author_home_url"))
        author_name = clean_optional(row.get("author_name")) or ""
        author_key = raw_author_id or (
            f"author_{stable_hash(platform, author_home_url)}"
            if author_home_url
            else f"author_{stable_hash(platform, author_name)}"
        )
        author_id = raw_author_id or author_key

        authors.append(
            {
                "author_id": author_id,
                "author_key": author_key,
                "platform": platform,
                "author_name": author_name,
                "author_home_url": author_home_url,
                "author_type": clean_optional(row.get("author_type")),
                "is_kol": to_bool(row.get("is_kol")),
                "fans_cnt": to_int(row.get("fans_cnt")),
                "author_desc": clean_optional(row.get("author_desc")),
                "ingest_batch_id": ingest_batch_id,
                "raw_source_key": "content_upload",
            }
        )
        engagement_total = (
            to_int(row.get("like_cnt"))
            + to_int(row.get("comment_cnt"))
            + to_int(row.get("share_cnt"))
            + to_int(row.get("favorite_cnt"))
        )
        contents.append(
            {
                "content_id": content_id,
                "content_key": content_key,
                "event_id": event_id,
                "author_id": author_id,
                "platform": platform,
                "source_url": source_url,
                "title": clean_optional(row.get("title")),
                "content_text": clean_optional(row.get("content_text")),
                "content_type": clean_optional(row.get("content_type")),
                "media_form": clean_optional(row.get("media_form")),
                "published_at": parse_time(row.get("published_at")),
                "like_cnt": to_int(row.get("like_cnt")),
                "comment_cnt": to_int(row.get("comment_cnt")),
                "share_cnt": to_int(row.get("share_cnt")),
                "favorite_cnt": to_int(row.get("favorite_cnt")),
                "view_cnt": to_int(row.get("view_cnt")),
                "engagement_total": engagement_total,
                "ingest_batch_id": ingest_batch_id,
                "raw_source_key": "content_upload",
            }
        )
        rel_event_content.append(
            {
                "event_id": event_id,
                "content_id": content_id,
                "match_type": "manual",
                "match_score": 1,
                "is_primary_event": True,
            }
        )
        rel_author_content.append(
            {
                "author_id": author_id,
                "content_id": content_id,
                "platform": platform,
                "published_at": parse_time(row.get("published_at")),
            }
        )

    return (
        pd.DataFrame(contents).drop_duplicates("content_key", keep="last"),
        pd.DataFrame(authors).drop_duplicates("author_key", keep="last"),
        pd.DataFrame(rel_event_content).drop_duplicates(["event_id", "content_id"], keep="last"),
        pd.DataFrame(rel_author_content).drop_duplicates(["author_id", "content_id"], keep="last"),
        pd.DataFrame(rejected),
    )


def standardize_comments(
    ods_comment: pd.DataFrame,
    dwd_content: pd.DataFrame,
    ingest_batch_id: str,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    content_by_id = {row["content_id"]: row for row in dwd_content.to_dict(orient="records")}
    content_by_source = {
        (normalize_text(row["platform"]).lower(), normalize_text(row["source_url"]).lower()): row
        for row in dwd_content.to_dict(orient="records")
    }
    rows: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    for row in ods_comment.to_dict(orient="records"):
        input_content_id = clean_optional(row.get("content_id"))
        platform = clean_optional(row.get("platform"))
        content_source_url = clean_optional(row.get("content_source_url"))
        content_row = content_by_id.get(input_content_id) if input_content_id else None
        if content_row is None and content_source_url:
            content_row = content_by_source.get(
                (normalize_text(platform).lower(), normalize_text(content_source_url).lower())
            )
            if content_row is None:
                content_row = next(
                    (
                        item
                        for (_platform, source_url), item in content_by_source.items()
                        if source_url == normalize_text(content_source_url).lower()
                    ),
                    None,
                )
        if content_row is None:
            rejected.append({"reason": "无法匹配内容", **row})
            continue
        content_id = content_row["content_id"]
        comment_author_name = clean_optional(row.get("comment_author_name")) or ""
        published_at = parse_time(row.get("published_at"))
        raw_comment_id = clean_optional(row.get("raw_comment_id"))
        comment_key = raw_comment_id or f"comment_{stable_hash(content_id, comment_author_name, row.get('comment_text'), published_at)}"
        comment_id = raw_comment_id or comment_key
        like_cnt = to_int(row.get("like_cnt"))
        reply_cnt = to_int(row.get("reply_cnt"))
        rows.append(
            {
                "comment_id": comment_id,
                "comment_key": comment_key,
                "content_id": content_id,
                "platform": platform or content_row.get("platform"),
                "comment_author_id": clean_optional(row.get("comment_author_id")),
                "comment_author_name": comment_author_name,
                "parent_comment_id": clean_optional(row.get("parent_comment_id")),
                "reply_level": to_int(row.get("reply_level"), 1),
                "comment_text": clean_optional(row.get("comment_text")),
                "published_at": published_at,
                "like_cnt": like_cnt,
                "reply_cnt": reply_cnt,
                "interaction_cnt": like_cnt + reply_cnt,
                "source_url": clean_optional(row.get("source_url")),
                "ingest_batch_id": ingest_batch_id,
                "raw_source_key": "comment_upload",
            }
        )
    return pd.DataFrame(rows).drop_duplicates("comment_key", keep="last"), pd.DataFrame(rejected)


def build_ads(
    dwd_event: pd.DataFrame,
    dwd_content: pd.DataFrame,
    dwd_author: pd.DataFrame,
    dwd_comment: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    content_enriched = dwd_content.copy()
    if not dwd_author.empty and not content_enriched.empty:
        content_enriched = content_enriched.merge(
            dwd_author[["author_id", "author_name", "is_kol"]],
            on="author_id",
            how="left",
        )
    if "author_name" not in content_enriched.columns:
        content_enriched["author_name"] = None
    if "is_kol" not in content_enriched.columns:
        content_enriched["is_kol"] = False

    overview_rows: list[dict[str, Any]] = []
    for event in dwd_event.to_dict(orient="records"):
        event_id = event["event_id"]
        event_contents = content_enriched[content_enriched["event_id"] == event_id]
        event_comments = dwd_comment[dwd_comment["content_id"].isin(event_contents["content_id"])]
        platform_counts = (
            event_contents.groupby("platform")["content_id"].nunique().sort_values(ascending=False)
            if not event_contents.empty
            else pd.Series(dtype="int64")
        )
        author_counts = (
            event_contents.groupby("author_name")["content_id"].nunique().sort_values(ascending=False)
            if not event_contents.empty
            else pd.Series(dtype="int64")
        )
        overview_rows.append(
            {
                "event_id": event_id,
                "event_name": event.get("event_name"),
                "event_type": event.get("event_type"),
                "brand_name": event.get("brand_name"),
                "model_name": event.get("model_name"),
                "start_time": event.get("start_time"),
                "end_time": event.get("end_time"),
                "event_status": event.get("event_status"),
                "platform_list": json.dumps(sorted(event_contents["platform"].dropna().unique().tolist()), ensure_ascii=False),
                "content_cnt": int(event_contents["content_id"].nunique()),
                "comment_cnt": int(event_comments["comment_id"].nunique()),
                "author_cnt": int(event_contents["author_id"].nunique()),
                "kol_content_cnt": int(event_contents[event_contents["is_kol"].fillna(False)]["content_id"].nunique()),
                "total_engagement": int(event_contents["engagement_total"].fillna(0).sum()),
                "top_platform_json": json.dumps(
                    [{"label": key, "value": int(value)} for key, value in platform_counts.head(10).items()],
                    ensure_ascii=False,
                ),
                "top_author_json": json.dumps(
                    [{"label": key, "value": int(value)} for key, value in author_counts.head(10).items()],
                    ensure_ascii=False,
                ),
                "data_lineage_json": json.dumps(
                    {"source": ["dwd_event", "dwd_content", "dwd_author", "dwd_comment"]},
                    ensure_ascii=False,
                ),
            }
        )

    trend_parts: list[pd.DataFrame] = []
    if not dwd_content.empty:
        content_daily = dwd_content.copy()
        content_daily["stat_date"] = pd.to_datetime(content_daily["published_at"], errors="coerce").dt.date
        content_daily = content_daily.dropna(subset=["stat_date"])
        trend_parts.append(
            content_daily.groupby(["event_id", "stat_date"], as_index=False).agg(
                content_cnt=("content_id", "nunique"),
                engagement_content=("engagement_total", "sum"),
            )
        )
    if not dwd_comment.empty and not dwd_content.empty:
        comment_daily = dwd_comment.merge(dwd_content[["content_id", "event_id"]], on="content_id", how="left")
        comment_daily["stat_date"] = pd.to_datetime(comment_daily["published_at"], errors="coerce").dt.date
        comment_daily = comment_daily.dropna(subset=["event_id", "stat_date"])
        trend_parts.append(
            comment_daily.groupby(["event_id", "stat_date"], as_index=False).agg(
                comment_cnt=("comment_id", "nunique"),
                engagement_comment=("interaction_cnt", "sum"),
            )
        )
    if trend_parts:
        trend = trend_parts[0]
        for part in trend_parts[1:]:
            trend = trend.merge(part, on=["event_id", "stat_date"], how="outer")
        for column in ["content_cnt", "comment_cnt", "engagement_content", "engagement_comment"]:
            if column not in trend.columns:
                trend[column] = 0
        trend = trend.fillna(0)
        trend["engagement_cnt"] = trend["engagement_content"] + trend["engagement_comment"]
        trend["data_lineage_json"] = json.dumps({"source": ["dwd_content", "dwd_comment"]}, ensure_ascii=False)
        trend = trend[["event_id", "stat_date", "content_cnt", "comment_cnt", "engagement_cnt", "data_lineage_json"]]
    else:
        trend = pd.DataFrame(columns=["event_id", "stat_date", "content_cnt", "comment_cnt", "engagement_cnt", "data_lineage_json"])

    rank_rows: list[dict[str, Any]] = []
    for event_id, event_contents in content_enriched.groupby("event_id"):
        for rank_type, frame in [
            ("engagement", event_contents.sort_values("engagement_total", ascending=False)),
            ("comment", event_contents.sort_values("comment_cnt", ascending=False)),
            ("kol_content", event_contents[event_contents["is_kol"].fillna(False)].sort_values("engagement_total", ascending=False)),
        ]:
            for rank_no, row in enumerate(frame.head(20).to_dict(orient="records"), start=1):
                rank_rows.append(
                    {
                        "rank_id": f"{event_id}_{rank_type}_{row['content_id']}",
                        "event_id": event_id,
                        "content_id": row["content_id"],
                        "rank_type": rank_type,
                        "rank_no": rank_no,
                        "platform": row.get("platform"),
                        "title": row.get("title"),
                        "author_name": row.get("author_name"),
                        "is_kol": bool(row.get("is_kol")),
                        "published_at": row.get("published_at"),
                        "engagement_total": row.get("engagement_total"),
                        "comment_cnt": row.get("comment_cnt"),
                        "source_url": row.get("source_url"),
                        "data_lineage_json": json.dumps({"source": ["dwd_content", "dwd_author"]}, ensure_ascii=False),
                    }
                )
    rank = pd.DataFrame(rank_rows)
    return pd.DataFrame(overview_rows), trend, rank


def run_etl(input_dir: Path, output_dir: Path) -> dict[str, int]:
    ingest_batch_id = f"batch_{pd.Timestamp.now().strftime('%Y%m%d_%H%M%S')}"
    output_dir.mkdir(parents=True, exist_ok=True)

    ods_event = read_upload(input_dir, EVENT_SPEC)
    ods_content = read_upload(input_dir, CONTENT_SPEC)
    ods_comment = read_upload(input_dir, COMMENT_SPEC)

    errors = []
    errors.extend(validate_required(ods_event, EVENT_SPEC))
    errors.extend(validate_required(ods_content, CONTENT_SPEC))
    errors.extend(validate_required(ods_comment, COMMENT_SPEC))
    if errors:
        raise ValueError("\n".join(errors))

    dwd_event = standardize_events(ods_event, ingest_batch_id)
    dwd_content, dwd_author, rel_event_content, rel_author_content, rejected_content = standardize_content(
        ods_content,
        dwd_event,
        ingest_batch_id,
    )
    dwd_comment, rejected_comment = standardize_comments(ods_comment, dwd_content, ingest_batch_id)
    ads_overview, ads_trend_daily, ads_content_rank = build_ads(dwd_event, dwd_content, dwd_author, dwd_comment)

    outputs = {
        "ods_event_upload": ods_event,
        "ods_content_upload": ods_content,
        "ods_comment_upload": ods_comment,
        "dwd_event": dwd_event,
        "dwd_content": dwd_content,
        "dwd_author": dwd_author,
        "dwd_comment": dwd_comment,
        "rel_event_content": rel_event_content,
        "rel_author_content": rel_author_content,
        "ads_event_overview": ads_overview,
        "ads_event_trend_daily": ads_trend_daily,
        "ads_event_content_rank": ads_content_rank,
        "rejected_content": rejected_content,
        "rejected_comment": rejected_comment,
    }
    for name, dataframe in outputs.items():
        dataframe.to_csv(output_dir / f"{name}.csv", index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(output_dir / "event_voc_etl_result.xlsx", engine="openpyxl") as writer:
        for name, dataframe in outputs.items():
            dataframe.to_excel(writer, sheet_name=name[:31], index=False)

    summary = {name: int(len(dataframe)) for name, dataframe in outputs.items()}
    (output_dir / "etl_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate VOC event templates and run local ODS->ADS ETL.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    template_parser = subparsers.add_parser("generate-templates")
    template_parser.add_argument("--output-dir", type=Path, default=DEFAULT_TEMPLATE_DIR)

    sample_parser = subparsers.add_parser("write-sample")
    sample_parser.add_argument("--output-dir", type=Path, default=DEFAULT_SAMPLE_DIR / "input")

    run_parser = subparsers.add_parser("run")
    run_parser.add_argument("--input-dir", type=Path, required=True)
    run_parser.add_argument("--output-dir", type=Path, required=True)

    demo_parser = subparsers.add_parser("run-sample")
    demo_parser.add_argument("--base-dir", type=Path, default=DEFAULT_SAMPLE_DIR)

    args = parser.parse_args()
    if args.command == "generate-templates":
        generate_templates(args.output_dir)
        print(f"generated templates: {args.output_dir}")
    elif args.command == "write-sample":
        write_sample_inputs(args.output_dir)
        print(f"generated sample inputs: {args.output_dir}")
    elif args.command == "run":
        summary = run_etl(args.input_dir, args.output_dir)
        print(json.dumps(summary, ensure_ascii=False, indent=2))
    elif args.command == "run-sample":
        input_dir = args.base_dir / "input"
        output_dir = args.base_dir / "output"
        write_sample_inputs(input_dir)
        summary = run_etl(input_dir, output_dir)
        print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
